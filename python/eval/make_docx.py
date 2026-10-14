"""Word (.docx) versions of the paper, the supplement and the response, laid out like the PDFs.

    cd python && python -m eval.make_docx          (after python -m eval.build_paper)

The Word files are meant to be *edited*, so they are built as native Word documents rather than as
pictures of the PDF:

* the IEEE Computer Society layout of the PDF: Letter paper, the same margins, a one-column title
  block above two 252 pt columns, Palatino body text and Arial headings and captions at the PDF's
  sizes, the running header with a live page number, the drop cap, the first-page author note;
* every figure and table is a floating text box placed where LaTeX placed it in the PDF (page,
  column and height are read from the built PDF), anchored to a paragraph on that page, so the
  pages break where the PDF's do; the box moves with its paragraph when text is edited;
* section headings are a Word multilevel list, captions and theorem heads are ``SEQ`` fields, and
  every cross-reference and citation is a ``REF`` field to a bookmark, so inserting a section,
  table, theorem or reference renumbers everything on *Update Fields* (F9);
* equations are native Word equations; tables are Word tables ruled like booktabs.

The text is taken from the same LaTeX sources and everything that carries a number is resolved from
the *built* paper (``main.aux``, ``supplement.aux``, ``main.bbl``, ``tables/numbers.tex``), so the
Word files and the PDFs cannot disagree. When Microsoft Word is installed, a final pass opens each
file in Word, updates every field and fits every text box, and saves it as a Word-native file
(``--no-word`` skips it).

It refuses to run on a build older than its sources, for the same reason ``make_dist`` does.
"""

from __future__ import annotations

import copy
import re
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pypandoc
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from eval.lab import corpus as C

PAPER = C.repo_root() / "paper"
DIST = PAPER / "dist"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
EMU = 12700  # per point
TW = 20  # twips per point

BODY_FONT = "Palatino Linotype"
SANS_FONT = "Arial"
MONO_FONT = "Courier New"

# IEEEtran compsoc journal geometry, measured from the built PDF (points).
PAGE_W, PAGE_H = 612, 792
MARGIN_X, MARGIN_TOP, MARGIN_BOTTOM, HEADER_DIST = 48, 44, 44, 26
COL_W, COL_GAP = 252, 12
TEXT_W = 2 * COL_W + COL_GAP
FLOAT_GAP = 9  # a vertical gap wider than this separates two floats on one page

PREAMBLE = r"""
\documentclass{article}
\usepackage{amsmath,amssymb,amsthm}
\newtheorem{theorem}{Theorem}
\newtheorem{lemma}{Lemma}
\newtheorem{corollary}{Corollary}
\newtheorem{proposition}{Proposition}
\theoremstyle{definition}
\newtheorem{definition}{Definition}
\newcommand{\advhist}{\mathrm{Adv}_{\mathrm{hist}}}
\newcommand{\advpost}{\mathrm{Adv}_{\mathrm{post}}}
\newcommand{\AdvH}{\mathsf{Adv}^{\mathrm{H}}}
\newcommand{\AdvR}{\mathsf{Adv}^{\mathrm{R}}}
\newcommand{\negl}{\mathsf{negl}}
\newcommand{\pt}[1]{\par QXPT0QX\textbf{#1}\par}
\newcommand{\resp}{\textit{Response.}\ }
\newcommand{\done}{QXDONE0QX}
\newcommand{\pdone}{QXPDONE0QX}
\newcommand{\planned}{QXPLAN0QX}
"""

TYPE_NAMES = {
    "section": ("Section", "Sections"),
    "subsection": ("Section", "Sections"),
    "subsubsection": ("Section", "Sections"),
    "appendix": ("Appendix", "Appendices"),
    "table": ("Table", "Tables"),
    "figure": ("Figure", "Figures"),
    "theorem": ("Theorem", "Theorems"),
    "lemma": ("Lemma", "Lemmas"),
    "proposition": ("Proposition", "Propositions"),
    "corollary": ("Corollary", "Corollaries"),
    "definition": ("Definition", "Definitions"),
    "equation": ("Eq.", "Eqs."),
    "enumi": ("item", "items"),
}
THEOREMS = ("theorem", "lemma", "proposition", "corollary", "definition")
TOKEN = re.compile(r"QX([A-Z]+?)(\d+)QX")
PARA_TOKENS = {
    "TITLE", "AUTHOR", "ABS", "KW", "SEP", "HL", "NONUM", "NI", "RUN", "DROP", "BIB", "BIO",
    "CAP", "FB", "FE", "RT", "RS", "PT",
}  # fmt: skip


# ----------------------------------------------------------------------------------- LaTeX side


def w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def labels_from(aux: Path) -> dict[str, tuple[str, str]]:
    """label -> (cleveref type, number as printed), from a built .aux file."""
    out: dict[str, tuple[str, str]] = {}
    if not aux.exists():
        return out
    text = aux.read_text(encoding="latin-1")
    for m in re.finditer(
        r"\\newlabel\{([^}]+)@cref\}\{\{\[([^\]]*)\]\[[^\]]*\]\[[^\]]*\]([^}]*)\}", text
    ):
        out[m.group(1)] = (m.group(2), m.group(3))
    return out


def pages_from(aux: Path) -> dict[str, int]:
    """label -> the PDF page it was typeset on."""
    text = aux.read_text(encoding="latin-1") if aux.exists() else ""
    return {
        m.group(1): int(m.group(2))
        for m in re.finditer(r"\\newlabel\{([^}@]+)\}\{\{[^}]*\}\{(\d+)\}", text)
    }


def citations_from(aux: Path) -> dict[str, str]:
    text = aux.read_text(encoding="latin-1") if aux.exists() else ""
    return dict(re.findall(r"\\bibcite\{([^}]+)\}\{([^}]+)\}", text))


def inline_inputs(text: str, seen: int = 0) -> str:
    if seen > 10:
        raise SystemExit("\\input nesting too deep")

    def rep(m: re.Match[str]) -> str:
        name = m.group(1)
        if name == "tables/numbers":
            return ""  # macros go in the preamble
        path = PAPER / (name if name.endswith(".tex") else name + ".tex")
        return inline_inputs(path.read_text(encoding="utf-8"), seen + 1)

    return re.sub(r"\\input\{([^}]+)\}", rep, text)


def strip_comments(text: str) -> str:
    return "\n".join(re.sub(r"(?<!\\)%.*$", "", line) for line in text.split("\n"))


def brace_arg(s: str, i: int) -> tuple[str, int]:
    """The balanced {...} group starting at s[i] == '{': (contents, index after it)."""
    if s[i] != "{":
        raise ValueError(f"expected '{{' at {i}: {s[i : i + 30]!r}")
    depth = 0
    for j in range(i, len(s)):
        if s[j] == "\\":
            continue
        if s[j] == "{" and s[j - 1] != "\\":
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1 : j], j + 1
    raise ValueError("unbalanced braces")


def num_key(n: str) -> tuple:
    return tuple(int(p) if p.isdigit() else p for p in re.split(r"[.]", n))


def consecutive(a: str, b: str) -> bool:
    ka, kb = num_key(a), num_key(b)
    return (
        len(ka) == len(kb)
        and ka[:-1] == kb[:-1]
        and isinstance(ka[-1], int)
        and isinstance(kb[-1], int)
        and kb[-1] == ka[-1] + 1
    )


def column_spec(inner: str, width: float) -> list[float | None]:
    """The columns of a float's tabular: a p{} width in points, or None for l, c and r."""
    m = re.search(r"\\begin\{tabular\*?\}(?:\{[^}]*\})?\{", inner)
    if not m:
        return []
    try:
        spec, _ = brace_arg(inner, m.end() - 1)
    except ValueError:
        return []
    spec = re.sub(r"@\{[^}]*\}|\|", "", spec)
    units = {"textwidth": TEXT_W, "columnwidth": COL_W, "linewidth": width}
    lengths = {"cm": 28.45, "mm": 2.845, "in": 72.27, "pt": 1.0}
    out: list[float | None] = []
    for c in re.finditer(r"[pmb]\{([^}]*)\}|[lcr]", spec):
        v = (c.group(1) or "").strip()
        if c.group(1) is None:
            out.append(None)
        elif mm := re.match(r"([0-9.]*)\\(textwidth|columnwidth|linewidth)", v):
            out.append(float(mm.group(1) or 1) * units[mm.group(2)])
        elif mm := re.match(r"([0-9.]+)\s*(cm|mm|in|pt)", v):
            out.append(float(mm.group(1)) * lengths[mm.group(2)])
        else:
            out.append(None)
    return out


def join_and(parts: list[str]) -> str:
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return f"{parts[0]} and {parts[1]}"
    return ", ".join(parts[:-1]) + " and " + parts[-1]


@dataclass
class Float:
    kind: str  # "table" | "figure"
    wide: bool
    label: str | None
    number: str
    page: int | None
    widths: list[float]  # each \includegraphics width, as a fraction of the float's width
    nowrap: bool = True  # no p{} column: LaTeX never wraps a cell, so Word must not either
    size: float = 8.0  # the table's font size, from \scriptsize, \footnotesize or \small
    colsep: float = 6.0  # LaTeX's \tabcolsep, the padding each side of a cell
    # one entry per column: a p{} column's width in points, or None for l, c and r
    columns: list = field(default_factory=list)


@dataclass
class Tokens:
    """What each QX<KIND><n>QX token in the flattened LaTeX stands for."""

    items: list[tuple] = field(default_factory=list)

    def add(self, kind: str, *data: object) -> str:
        self.items.append((kind, *data))
        return f"QX{kind}{len(self.items) - 1}QX"


class Source:
    """One LaTeX document flattened for pandoc, with tokens where Word needs fields or layout."""

    def __init__(self, root: str, labels: dict, pages: dict, cites: dict, fields: bool):
        self.root, self.labels, self.pages, self.cites = root, labels, pages, cites
        self.fields = fields
        self.tok = Tokens()
        self.floats: list[Float] = []

    # References: cleveref's \cref (sorted, grouped by type, runs of three or more as "a to b") and
    # the cite package's \cite (sorted, runs of three or more as "[a]--[b]").
    def ref_token(self, key: str) -> str:
        if key not in self.labels:
            raise SystemExit(f"unresolved reference: {key}")
        return self.tok.add("REF", key, self.labels[key][1])

    def cref(self, keys: list[str], capital: bool = True) -> str:
        groups: dict[tuple[str, str], list[tuple[str, str]]] = {}
        for k in keys:
            if k not in self.labels:
                raise SystemExit(f"unresolved reference: {k}")
            typ, num = self.labels[k]
            names = TYPE_NAMES.get(typ, (typ.capitalize(), typ.capitalize() + "s"))
            if typ == "figure" and not capital:  # IEEEtran's cleveref: \cref gives "Fig."
                names = ("Fig.", "Figs.")
            groups.setdefault(names, []).append((k, num))
        out = []
        for (one, many), refs in groups.items():
            refs.sort(key=lambda r: num_key(r[1]))
            runs: list[list[tuple[str, str]]] = []
            for r in refs:
                if runs and consecutive(runs[-1][-1][1], r[1]):
                    runs[-1].append(r)
                else:
                    runs.append([r])
            parts = []
            for run in runs:
                if len(run) >= 3:
                    parts.append(f"{self.ref_token(run[0][0])} to {self.ref_token(run[-1][0])}")
                else:
                    parts.extend(self.ref_token(k) for k, _ in run)
            out.append(f"{one if len(refs) == 1 else many}~{join_and(parts)}")
        return join_and(out)

    def cite(self, keys: list[str]) -> str:
        bad = [k for k in keys if k not in self.cites]
        if bad:
            raise SystemExit(f"uncited keys: {bad}")
        nums = sorted({int(self.cites[k]): k for k in keys}.items())
        runs: list[list[tuple[int, str]]] = []
        for n, k in nums:
            if runs and runs[-1][-1][0] == n - 1:
                runs[-1].append((n, k))
            else:
                runs.append([(n, k)])
        parts = []
        for run in runs:
            if len(run) >= 3:
                parts.append(
                    self.tok.add("CIT", run[0][1], f"[{run[0][0]}]")
                    + "--"
                    + self.tok.add("CIT", run[-1][1], f"[{run[-1][0]}]")
                )
            else:
                parts.extend(self.tok.add("CIT", k, f"[{n}]") for n, k in run)
        return ", ".join(parts)

    def floats_pass(self, body: str) -> str:
        def rep(m: re.Match[str]) -> str:
            kind, star, inner = m.group(1), m.group(2), m.group(4)
            lab = re.search(r"\\label\{([^}]+)\}", inner)
            label = lab.group(1) if lab else None
            number = self.labels.get(label, ("", ""))[1] if label else ""
            widths = []
            for g in re.finditer(r"\\includegraphics(?:\[([^\]]*)\])?", inner):
                wm = re.search(
                    r"width=([0-9.]*)\\(linewidth|columnwidth|textwidth)", g.group(1) or ""
                )
                frac = float(wm.group(1) or 1) if wm else 1.0
                if wm and wm.group(2) == "textwidth" and not star:
                    frac *= TEXT_W / COL_W
                widths.append(min(frac, 1.0))
            fid = len(self.floats)
            self.floats.append(
                Float(
                    kind,
                    bool(star),
                    label,
                    number,
                    self.pages.get(label or ""),
                    widths,
                    size=next(
                        (
                            v
                            for k, v in (("scriptsize", 7.0), ("footnotesize", 8.0), ("small", 9.0))
                            if "\\" + k in inner.split("\\begin{tabular")[0]
                        ),
                        9.5 if kind == "table" else 8.0,
                    ),
                    colsep=float(cs.group(1))
                    if (cs := re.search(r"tabcolsep\}\{([0-9.]+)pt", inner))
                    else 6.0,
                    columns=(cols := column_spec(inner, TEXT_W if star else COL_W)),
                    nowrap=all(c is None for c in cols),
                )
            )
            inner = inner.replace("\\caption{", "\\caption{" + self.tok.add("CAP", fid) + " ", 1)
            return (
                f"\n\n{self.tok.add('FB', fid)}\n\n\\begin{{{kind}}}{inner}\\end{{{kind}}}"
                f"\n\n{self.tok.add('FE', fid)}\n\n"
            )

        return re.sub(
            r"\\begin\{(table|figure)(\*?)\}(\[[^\]]*\])?(.*?)\\end\{\1\2\}", rep, body, flags=re.S
        )

    def headings_pass(self, body: str) -> str:
        out, i = [], 0
        pat = re.compile(r"\\(section|subsection|subsubsection)(\*?)\{")
        while m := pat.search(body, i):
            out.append(body[i : m.start()])
            title, j = brace_arg(body, m.end() - 1)
            labels = []
            while lm := re.match(r"\s*\\label\{([^}]+)\}", body[j:]):
                labels.append(lm.group(1))
                j += lm.end()
            mark = self.tok.add("NONUM") if m.group(2) else ""
            if labels:
                mark += self.tok.add("HL", labels)
            out.append(f"\\{m.group(1)}{m.group(2)}{{{mark} {title}}}")
            i = j
        out.append(body[i:])
        return "".join(out)

    def refs_pass(self, body: str) -> str:
        if self.fields:
            body = re.sub(
                r"\\([cC])ref\{([^}]+)\}",
                lambda m: self.cref(
                    [k.strip() for k in m.group(2).split(",") if k.strip()], m.group(1) == "C"
                ),
                body,
            )
            body = re.sub(r"\\ref\{([^}]+)\}", lambda m: self.ref_token(m.group(1)), body)
        else:  # documents whose targets live in another file: plain text
            body = re.sub(
                r"\\([cC])ref\{([^}]+)\}",
                lambda m: self.plain(
                    self.cref([k.strip() for k in m.group(2).split(",")], m.group(1) == "C")
                ),
                body,
            )
            body = re.sub(
                r"\\ref\{([^}]+)\}", lambda m: self.plain(self.ref_token(m.group(1))), body
            )
        return re.sub(
            r"\\cite\{([^}]+)\}",
            lambda m: self.cite([k.strip() for k in m.group(1).split(",")]),
            body,
        )

    def plain(self, s: str) -> str:
        return TOKEN.sub(lambda m: self.tok.items[int(m.group(2))][2], s)

    def bibliography(self, bbl: Path) -> str:
        text = bbl.read_text(encoding="utf-8")
        body = text[text.index("\\bibitem") :]
        body = body[: body.index("\\end{thebibliography}")]
        items = re.split(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", body)[1:]
        out = [f"\\section*{{{self.tok.add('NONUM')} References}}"]
        for key, entry in zip(items[0::2], items[1::2], strict=True):
            e = re.sub(r"\\BIBentry(?:STD|ALT)interwordspacing", "", entry)
            e = e.replace("\\newblock", "")
            e = re.sub(r"\\bibinfo\{[^}]*\}", "", e)
            e = re.sub(r"\\url\{([^}]*)\}", r"\\url{\1}", e)
            e = re.sub(r"\\href\{[^}]*\}\{([^}]*)\}", r"\1", e)
            out.append(f"{self.tok.add('BIB', key)} {' '.join(e.split())}\n")
        return "\n\n".join(out)

    def normalise(self, body: str) -> str:
        body = re.sub(r"\\texorpdfstring\{([^}]*)\}\{[^}]*\}", r"\1", body)
        body = re.sub(r"\\nolinkurl\{([^}]*)\}", r"\\texttt{\1}", body)
        body = re.sub(r"\\bstctlcite\{[^}]*\}", "", body)
        for cmd in ("\\IEEEpeerreviewmaketitle", "\\IEEEdisplaynontitleabstractindextext"):
            body = body.replace(cmd, "")
        body = body.replace("\\maketitle", "")
        m = re.search(r"\\IEEEPARstart\{([^}]*)\}\{([^}]*)\}", body)
        if m:
            body = (
                body[: m.start()]
                + f"{self.tok.add('DROP', m.group(1))}\\textsc{{{m.group(2)}}}"
                + body[m.end() :]
            )
        for sym, ch in (("\\S", "§"), ("\\ddagger", "‡"), ("\\dagger", "†"), ("\\ast", "*")):
            body = body.replace(f"$^{sym}$", f"\\textsuperscript{{{ch}}}")
        while (i := body.find("\\paragraph{")) >= 0:
            head, j = brace_arg(body, i + len("\\paragraph"))
            body = body[:i] + f"\n\n{self.tok.add('RUN')} {head}: " + body[j:].lstrip()
        body = re.sub(r"\\noindent\s*", lambda m: self.tok.add("NI") + " ", body)
        body = re.sub(r"(\\includegraphics(?:\[[^\]]*\])?\{[^}]*?)\.pdf\}", r"\1.png}", body)
        body = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}", "", body)
        for a, b in (
            (r"\bigl|", r"\left|"),
            (r"\bigr|", r"\right|"),
            (r"\Bigl(", r"\left("),
            (r"\Bigr)", r"\right)"),
            (r"\bigl(", r"\left("),
            (r"\bigr)", r"\right)"),
            (r"\tfrac", r"\frac"),
        ):
            body = body.replace(a, b)
        body = re.sub(
            r"\\begin\{IEEEbiographynophoto\}\{([^}]*)\}(.*?)\\end\{IEEEbiographynophoto\}",
            lambda m: (
                f"\n\n{self.tok.add('BIO')} \\textbf{{{m.group(1)}}} {m.group(2).strip()}\n\n"
            ),
            body,
            flags=re.S,
        )
        body = re.sub(r"\$\\checkmark\$|\\checkmark\b", "\u2713", body)
        body = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", body)  # pandoc prints these as text
        for cmd in ("\\vfill", "\\clearpage", "\\newpage"):
            body = body.replace(cmd, "")
        return body


def ieee_front(src: str, s: Source, abstract: str | None, keywords: str | None) -> str:
    title, _ = brace_arg(src, src.index("\\title{") + len("\\title"))
    author_block, _ = brace_arg(src, src.index("\\author{") + len("\\author"))
    name = author_block.split("%")[0].split("\\IEEEcompsocitemizethanks")[0].strip()
    thanks = re.search(r"\\IEEEcompsocthanksitem\s*(.*?)\}\s*$", author_block, re.S)
    note = (thanks.group(1) if thanks else "").replace("\\protect\\\\", "\\\\").strip()
    front = f"{s.tok.add('TITLE')} {title}\n\n"
    front += f"{s.tok.add('AUTHOR')} {name}" + (f"\\footnote{{{note}}}" if note else "") + "\n\n"
    if abstract:
        front += f"{s.tok.add('ABS')} \\textbf{{Abstract}}---{abstract.strip()}\n\n"
    if keywords:
        front += f"{s.tok.add('KW')} \\textbf{{Index Terms}}---{keywords}\n\n"
    return front + f"{s.tok.add('SEP')}\n\n"


def flatten(root: str, fields: bool) -> tuple[str, Source]:
    tex = PAPER / f"{root}.tex"
    src = strip_comments(tex.read_text(encoding="utf-8"))
    full = inline_inputs(src)
    start = full.index("\\begin{document}") + len("\\begin{document}")
    body = full[start : full.index("\\end{document}")]
    labels = labels_from(PAPER / f"{root}.aux")
    if root != "main":
        labels = {**labels_from(PAPER / "main.aux"), **labels}
    if root == "response":  # \externaldocument[supp-]{supplement}
        labels.update({"supp-" + k: v for k, v in labels_from(PAPER / "supplement.aux").items()})
    s = Source(
        root, labels, pages_from(PAPER / f"{root}.aux"), citations_from(PAPER / "main.aux"), fields
    )

    if root == "response":
        title, _ = brace_arg(src, src.index("\\title{") + len("\\title"))
        head, _, sub = title.partition("\\\\")
        sub = sub.replace("\\large", "").strip()
        front = f"{s.tok.add('RT')} {head.strip()}\n\n{s.tok.add('RS')} {sub}\n\n"
        body = front + body.replace("\\maketitle", "", 1)
    else:
        for marker in ("\\IEEEpeerreviewmaketitle", "\\maketitle"):
            if marker in body:
                body = body[body.index(marker) + len(marker) :]
                break
        abstract = keywords = None
        if root == "main":
            abstract = strip_comments((PAPER / "sections" / "00-abstract.tex").read_text("utf-8"))
            keywords = " ".join(
                re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", src, re.S)
                .group(1)
                .split()
            )
        body = ieee_front(src, s, abstract, keywords) + body

    body = s.floats_pass(body)
    body = s.headings_pass(body)
    if "\\bibliography{refs}" in body:
        body = body.replace("\\bibliographystyle{IEEEtran}", "")
        body = body.replace("\\bibliography{refs}", s.bibliography(PAPER / "main.bbl"))
    body = s.refs_pass(body)
    body = s.normalise(body)
    numbers = (PAPER / "tables" / "numbers.tex").read_text(encoding="utf-8")
    if "[not run]" in numbers:
        sys.exit("numbers.tex has unmeasured values; run the experiments first")
    return PREAMBLE + numbers + "\n\\begin{document}\n" + body + "\n\\end{document}\n", s


# ------------------------------------------------------------------------------------ PDF side


@dataclass
class Box:
    page: int
    x: float
    y: float
    w: float
    h: float
    float_page: bool = False  # LaTeX gave this float a page of floats, with no text on it


def norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKC", text).lower())


class PdfLayout:
    """Where LaTeX put each float, and the text of each page, read from the built PDF."""

    def __init__(self, pdf: Path):
        import pymupdf

        self.doc = pymupdf.open(str(pdf))
        self.lines: list[list[tuple]] = []
        self.text: list[str] = []
        for page in self.doc:
            lines = []
            for b in page.get_text("dict")["blocks"]:
                for ln in b.get("lines", []):
                    spans = [sp for sp in ln["spans"] if sp["text"].strip()]
                    if not spans:
                        continue
                    t = "".join(sp["text"] for sp in ln["spans"])
                    body = sum(len(sp["text"]) for sp in spans if self.is_body(sp))
                    lines.append((*ln["bbox"], t, body * 2 > sum(len(sp["text"]) for sp in spans)))
            self.lines.append(lines)
            self.text.append(norm(" ".join(ln[4] for ln in lines if ln[1] > 40)))
        # The running text alone, in reading order (left column, then right), as one stream with
        # the offset at which each page begins: paragraphs are located in it in sequence.
        self.stream, self.page_at = "", []
        for lines in self.lines:
            self.page_at.append(len(self.stream))
            body = sorted(
                (ln for ln in lines if ln[5] and ln[1] > 40), key=lambda ln: (ln[0] > 306, ln[1])
            )
            self.stream += norm(" ".join(ln[4] for ln in body))

    def page_of(self, offset: int) -> int:
        return max(i for i, start in enumerate(self.page_at) if start <= offset) + 1

    @staticmethod
    def is_body(sp: dict) -> bool:
        f, size = sp["font"], sp["size"]
        return (
            ("Palladio" in f and size > 9.0)
            or ("NimbusSan" in f and "Bold" in f and size >= 9.4)
            or (f.startswith("CM") and size >= 9.0)
        )

    def float_box(self, fl: Float) -> Box | None:
        if not fl.page or not fl.number or fl.page > len(self.lines):
            return None
        page = self.doc[fl.page - 1]
        lines = self.lines[fl.page - 1]
        want = f"TABLE {fl.number}" if fl.kind == "table" else f"Fig. {fl.number}."
        cap = next((ln for ln in lines if ln[4].strip().startswith(want)), None)
        if cap is None:
            return None
        centre = (cap[0] + cap[2]) / 2
        if centre < 240:
            x0, x1 = MARGIN_X, MARGIN_X + COL_W
        elif centre > 372:
            x0, x1 = MARGIN_X + COL_W + COL_GAP, PAGE_W - MARGIN_X
        else:
            x0, x1 = MARGIN_X, PAGE_W - MARGIN_X

        def in_x(b: tuple) -> bool:
            return b[2] > x0 + 2 and b[0] < x1 - 2

        body = [ln for ln in lines if ln[5] and in_x(ln) and ln[1] > 40]
        top = max((ln[3] for ln in body if ln[3] <= cap[1] + 1), default=40)
        bottom = min((ln[1] for ln in body if ln[1] >= cap[3] - 1), default=PAGE_H)
        objs = [ln[:4] for ln in lines if not ln[5] and in_x(ln) and top < ln[1] and ln[3] < bottom]
        objs += [
            tuple(d["rect"])
            for d in page.get_drawings()
            if in_x(tuple(d["rect"])) and top < d["rect"][1] and d["rect"][3] < bottom
        ]
        objs += [
            tuple(i["bbox"])
            for i in page.get_image_info()
            if in_x(tuple(i["bbox"])) and top < i["bbox"][1] and i["bbox"][3] < bottom
        ]
        # Several floats can share the region (a page of floats): split it at the gaps LaTeX leaves
        # between floats, and keep the block that holds this float's caption.
        objs.sort(key=lambda o: o[1])
        blocks: list[list[float]] = []
        for o in objs:
            if blocks and o[1] <= blocks[-1][1] + FLOAT_GAP:
                blocks[-1][1] = max(blocks[-1][1], o[3])
            else:
                blocks.append([o[1], o[3]])
        caps = [
            ln for ln in lines if in_x(ln) and re.match(r"(TABLE \d+$|Fig\. \d+\.)", ln[4].strip())
        ]

        def caption_in(b: list[float]) -> str | None:
            c = next((ln for ln in caps if b[0] - 1 <= ln[1] <= b[1] + 1), None)
            return None if c is None else c[4].strip()

        # A table's caption heads its float and owns what follows; a figure's caption closes its
        # float and owns the picture just above it.
        groups: list[list[list[float]]] = []
        pending: list[list[float]] = []
        for blk in blocks:
            c = caption_in(blk)
            if c is None:
                pending.append(blk)
            elif c.startswith("TABLE"):
                if groups:
                    groups[-1].extend(pending)
                pending = []
                groups.append([blk])
            else:
                if groups and pending[:-1]:
                    groups[-1].extend(pending[:-1])
                groups.append([*pending[-1:], blk])
                pending = []
        if groups:
            groups[-1].extend(pending)
        mine = next(g for g in groups if any(b[0] - 1 <= cap[1] <= b[1] + 1 for b in g))
        y0, y1 = min(b[0] for b in mine), max(b[1] for b in mine)
        float_page = not any(ln[5] for ln in lines if ln[1] > 40)
        return Box(fl.page, x0, y0, x1 - x0, y1 - y0, float_page)


# ----------------------------------------------------------------------------------- Word side


def el(tag: str, **attrs: object) -> etree._Element:
    e = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def text_run(text: str, rpr: etree._Element | None = None) -> etree._Element:
    r = el("r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = el("t")
    t.text = text
    t.set(XML_SPACE, "preserve")
    r.append(t)
    return r


def field_runs(instr: str, result: str, rpr: etree._Element | None) -> list[etree._Element]:
    def fc(kind: str) -> etree._Element:
        r = el("r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        r.append(el("fldChar", fldCharType=kind))
        return r

    ins = el("r")
    if rpr is not None:
        ins.append(copy.deepcopy(rpr))
    it = el("instrText")
    it.text = f" {instr} "
    it.set(XML_SPACE, "preserve")
    ins.append(it)
    return [fc("begin"), ins, fc("separate"), text_run(result, rpr), fc("end")]


def ppr(p: etree._Element) -> etree._Element:
    return p.get_or_add_pPr()


def set_style(p: etree._Element, style_id: str) -> None:
    pp = ppr(p)
    st = pp.find(w("pStyle"))
    if st is None:
        st = el("pStyle")
        pp.insert(0, st)
    st.set(qn("w:val"), style_id)


def set_ind(p: etree._Element, **attrs: int) -> None:
    pp = ppr(p)
    ind = pp.find(w("ind"))
    if ind is None:
        ind = el("ind")
        pp.append(ind)
    for k, v in attrs.items():
        ind.set(qn(f"w:{k}"), str(v))


def set_num(p: etree._Element, num_id: int, ilvl: int = 0) -> None:
    pp = ppr(p)
    old = pp.find(w("numPr"))
    if old is not None:
        pp.remove(old)
    np_ = el("numPr")
    np_.append(el("ilvl", val=ilvl))
    np_.append(el("numId", val=num_id))
    st = pp.find(w("pStyle"))
    pp.insert(1 if st is not None else 0, np_)


def ordered_ppr(pp: etree._Element) -> None:
    """Put pPr children in schema order; Word rejects some out-of-order combinations."""
    order = [
        "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
        "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku",
        "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
        "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
        "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
        "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
    ]  # fmt: skip
    rank = {w(t): i for i, t in enumerate(order)}
    kids = sorted(pp, key=lambda e: rank.get(e.tag, 99))
    for k in kids:
        pp.remove(k)
        pp.append(k)


def para_text(p: etree._Element) -> str:
    return "".join(t.text or "" for t in p.iter(w("t")))


def lead_text(p: etree._Element) -> str:
    """The paragraph's text up to its first equation (the PDF spells equations differently)."""
    out = []
    for e in p.iter(w("t"), f"{{{M_NS}}}oMath", f"{{{M_NS}}}oMathPara"):
        if any(True for _ in e.iterancestors(w("txbxContent"))):
            continue  # a float anchored here is not part of the paragraph's text
        if e.tag != w("t"):
            break
        out.append(e.text or "")
    return "".join(out)


def merge_runs(p: etree._Element) -> None:
    prev = prev_t = prev_key = None
    for r in list(p.iterchildren(w("r"))):
        kids = [c for c in r if c.tag != w("rPr")]
        if len(kids) == 1 and kids[0].tag == w("t"):
            rpr = r.find(w("rPr"))
            key = etree.tostring(rpr) if rpr is not None else b""
            if prev is not None and key == prev_key and r.getprevious() is prev:
                prev_t.text = (prev_t.text or "") + (kids[0].text or "")
                prev_t.set(XML_SPACE, "preserve")
                p.remove(r)
                continue
            prev, prev_t, prev_key = r, kids[0], key
        else:
            prev = None


class Bookmarks:
    def __init__(self) -> None:
        self.n = 0
        self.by_label: dict[str, tuple[str, str]] = {}  # label -> (bookmark, "para" | "seq")

    def new(self) -> tuple[str, etree._Element, etree._Element]:
        self.n += 1
        name = f"_Ref{500000000 + self.n}"
        return (
            name,
            el("bookmarkStart", id=9000 + self.n, name=name),
            el("bookmarkEnd", id=9000 + self.n),
        )


@dataclass
class Layout:
    """How a document is laid out: the IEEE two-column paper, or a one-column letter."""

    ieee: bool
    header: str = ""


class Builder:
    def __init__(self, path: Path, src: Source, layout: Layout, pdf: PdfLayout | None):
        self.path, self.src, self.layout, self.pdf = path, src, layout, pdf
        self.d = docx.Document(str(path))
        self.body = self.d.element.body
        self.bm = Bookmarks()
        self.markers: dict[tuple[str, int], etree._Element] = {}
        self.docpr = 7000
        self.cursor = 0
        self.grids: dict[etree._Element, list[int]] = {}
        self.warnings: list[str] = []

    # --------------------------------------------------------------------- tokens and fields
    def run(self) -> None:
        theorem_paras = self.pandoc_bookmarks()
        for p in self.body.iter(w("p")):
            merge_runs(p)
        roles = self.paragraph_tokens()
        self.styles()
        self.numbering()
        for p, kind, idx in roles:
            self.apply_role(p, kind, idx)
        self.theorem_heads(theorem_paras)
        self.run_tokens()
        self.tables()
        self.figures_paragraphs()
        self.tall_math()
        self.lists()
        if self.layout.ieee:
            self.floats()
        self.sections()
        for p in self.body.iter(w("p")):
            merge_runs(p)
            if p.pPr is not None:
                ordered_ppr(p.pPr)
        self.settings()
        for e in self.d.element.iter():
            for k, v in e.attrib.items():
                if "QX" in v:
                    e.set(
                        k,
                        TOKEN.sub(
                            lambda m: (
                                str(self.src.tok.items[int(m.group(2))][-1])
                                if m.group(1) in ("REF", "CIT")
                                else ""
                            ),
                            v,
                        ).strip(),
                    )
        for s in self.d.styles.element.iter(w("pPr")):
            ordered_ppr(s)
        self.d.save(str(self.path))
        self.footnote_mark()

    def pandoc_bookmarks(self) -> dict[str, etree._Element]:
        """Theorem labels -> their paragraph; then drop pandoc's bookmarks (ours replace them)."""
        found = {}
        for bs in list(self.body.iter(w("bookmarkStart"))):
            name = bs.get(qn("w:name"))
            if self.src.labels.get(name, ("",))[0] in THEOREMS:
                p = bs.getparent() if bs.getparent().tag == w("p") else bs.getnext()
                while p is not None and p.tag != w("p"):
                    p = p.getnext()
                if p is not None:
                    found[name] = p
        for b in list(self.body.iter(w("bookmarkStart"), w("bookmarkEnd"))):
            b.getparent().remove(b)
        return found

    def paragraph_tokens(self) -> list[tuple]:
        roles: list[tuple] = []
        for p in list(self.body.iter(w("p"))):
            before = len(roles)
            for t in p.iter(w("t")):
                if not t.text or "QX" not in t.text:
                    continue

                def take(m: re.Match[str], p: etree._Element = p) -> str:
                    if m.group(1) in PARA_TOKENS:
                        roles.append((p, m.group(1), int(m.group(2))))
                        return ""
                    return m.group(0)

                t.text = TOKEN.sub(take, t.text)
            if len(roles) > before:  # drop the space a removed leading token leaves behind
                first = next((t for t in p.iter(w("t")) if t.text), None)
                if first is not None:
                    first.text = first.text.lstrip()
        return roles

    def apply_role(self, p: etree._Element, kind: str, idx: int) -> None:
        item = self.src.tok.items[idx]
        if kind in ("FB", "FE"):
            self.markers[(kind, item[1])] = p
        elif kind == "TITLE":
            set_style(p, "PaperTitle")
        elif kind == "AUTHOR":
            set_style(p, "PaperAuthor")
        elif kind == "ABS":
            set_style(p, "AbstractText")
        elif kind == "KW":
            set_style(p, "IndexTerms")
        elif kind == "SEP":
            self.separator(p)
        elif kind == "RT":
            set_style(p, "PaperTitle")
        elif kind == "RS":
            set_style(p, "Subtitle")
        elif kind == "PT":
            set_style(p, "PointHead")
        elif kind == "NONUM":
            set_num(p, 0)
        elif kind == "HL":
            name, bs, be = self.bm.new()
            self.wrap_paragraph(p, bs, be)
            for lab in item[1]:
                self.bm.by_label[lab] = (name, "para")
        elif kind == "NI":
            set_ind(p, firstLine=0)
        elif kind == "RUN":
            set_ind(p, firstLine=int(28.5 * TW))
        elif kind == "DROP":
            self.drop_cap(p, item[1])
        elif kind == "BIB":
            set_style(p, "Reference")
            set_num(p, 901)
            name, bs, be = self.bm.new()
            self.wrap_paragraph(p, bs, be)
            self.bm.by_label["bib:" + item[1]] = (name, "para")
        elif kind == "BIO":
            set_style(p, "Biography")
        elif kind == "CAP":
            self.caption(p, self.src.floats[item[1]])

    def wrap_paragraph(self, p: etree._Element, bs: etree._Element, be: etree._Element) -> None:
        pp = p.find(w("pPr"))
        p.insert(p.index(pp) + 1 if pp is not None else 0, bs)
        p.append(be)

    def seq(self, name: str, number: str, rpr: etree._Element | None) -> tuple[str, list]:
        bname, bs, be = self.bm.new()
        return bname, [bs, *field_runs(f"SEQ {name} \\* ARABIC", number, rpr), be]

    def caption(self, p: etree._Element, fl: Float) -> None:
        if fl.kind == "table":
            set_style(p, "TableCaption")
            label = el("p")
            label.append(el("pPr"))
            set_style(label, "TableLabel")
            label.append(text_run("TABLE "))
            bname, runs = self.seq("Table", fl.number, None)
            for r in runs:
                label.append(r)
            p.addprevious(label)
        else:
            set_style(p, "ImageCaption")
            first = next(iter(p.iterchildren(w("r"))), None)
            bname, runs = self.seq("Figure", fl.number, None)
            new = [text_run("Fig. "), *runs, text_run(". ")]
            if first is not None:
                for r in new:
                    first.addprevious(r)
            else:
                p.extend(new)
        if fl.label:
            self.bm.by_label[fl.label] = (bname, "seq")

    def theorem_heads(self, labelled: dict[str, etree._Element]) -> None:
        by_para = {id(p): lab for lab, p in labelled.items()}
        pat = re.compile(r"^(Theorem|Lemma|Proposition|Corollary|Definition) (\d+)$")
        for p in self.body.iter(w("p")):
            runs = list(p.iterchildren(w("r")))
            if (
                not runs
                or runs[0].find(w("rPr")) is None
                or runs[0].find(f"{w('rPr')}/{w('b')}") is None
            ):
                continue
            m = pat.match(para_text_of(runs[0]))
            if not m:
                continue
            rpr = runs[0].find(w("rPr"))
            bname, seq = self.seq(m.group(1), m.group(2), rpr)
            head = text_run(m.group(1) + " ", rpr)
            runs[0].addprevious(head)
            for r in seq:
                runs[0].addprevious(r)
            p.remove(runs[0])
            # the optional title, "(...).", is part of the head in IEEEtran
            after = [
                e for e in seq[-1].itersiblings() if e.tag == w("r") and e.find(w("t")) is not None
            ]
            if after and "".join(para_text_of(r) for r in after[:2]).lstrip().startswith("("):
                for e in after:
                    erpr = e.find(w("rPr"))
                    if erpr is None:
                        erpr = el("rPr")
                        e.insert(0, erpr)
                    if erpr.find(w("i")) is not None:
                        break
                    if erpr.find(w("b")) is None:
                        erpr.insert(0, el("b"))
                    if para_text_of(e).rstrip().endswith(")."):
                        break
            lab = by_para.get(id(p))
            if lab:
                want = self.src.labels[lab][1]
                if want != m.group(2):
                    self.warnings.append(f"{lab}: Word numbers it {m.group(2)}, LaTeX {want}")
                self.bm.by_label[lab] = (bname, "seq")

    def run_tokens(self) -> None:
        for r in list(self.body.iter(w("r"))):
            t = r.find(w("t"))
            if t is None or not t.text or "QX" not in t.text:
                continue
            rpr = r.find(w("rPr"))
            pieces: list[etree._Element] = []
            pos = 0
            for m in TOKEN.finditer(t.text):
                if m.start() > pos:
                    pieces.append(text_run(t.text[pos : m.start()], rpr))
                pieces.extend(self.token_runs(m.group(1), int(m.group(2)), rpr))
                pos = m.end()
            if pos < len(t.text):
                pieces.append(text_run(t.text[pos:], rpr))
            for piece in pieces:
                r.addprevious(piece)
            r.getparent().remove(r)

    def token_runs(self, kind: str, idx: int, rpr: etree._Element | None) -> list:
        item = self.src.tok.items[idx]
        if kind == "REF":
            target = self.bm.by_label.get(item[1])
            if target is None:
                return [text_run(item[2], rpr)]
            name, how = target
            sw = "\\w \\h \\* CHARFORMAT" if how == "para" else "\\h \\* CHARFORMAT"
            return field_runs(f"REF {name} {sw}", item[2], rpr)
        if kind == "CIT":
            target = self.bm.by_label.get("bib:" + item[1])
            if target is None:
                return [text_run(item[2], rpr)]
            return field_runs(f"REF {target[0]} \\w \\h \\* CHARFORMAT", item[2], rpr)
        if kind in ("DONE", "PDONE", "PLAN"):
            text, colour = {
                "DONE": ("[done]", "1E8C1E"),
                "PDONE": ("[partly done]", "CC6600"),
                "PLAN": ("[in progress]", "1A1AB3"),
            }[kind]
            new = copy.deepcopy(rpr) if rpr is not None else el("rPr")
            new.append(el("b"))
            new.append(el("color", val=colour))
            return [text_run(text, new)]
        self.warnings.append(f"unhandled token {kind}{idx}")
        return []

    # ------------------------------------------------------------------------------- layout
    def separator(self, p: etree._Element) -> None:
        """IEEEtran compsoc's rule with a diamond under the abstract."""
        for r in list(p.iterchildren(w("r"))):
            p.remove(r)
        set_style(p, "Separator")
        lift = el("rPr")
        lift.append(el("position", val=5))
        p.append(text_run("\t", lift))
        dia = el("rPr")
        rf = el("rFonts", ascii="Segoe UI Symbol", hAnsi="Segoe UI Symbol", cs="Segoe UI Symbol")
        dia.append(rf)
        dia.append(el("sz", val=14))
        p.append(text_run("\u25c6", dia))
        p.append(text_run("\t", lift))
        r = p[-1]
        r.find(w("t")).tag = w("tab")
        p[-3].find(w("t")).tag = w("tab")

    def drop_cap(self, p: etree._Element, letter: str) -> None:
        cap = parse_xml(
            f'<w:p xmlns:w="{W_NS}"><w:pPr><w:pStyle w:val="FirstParagraph"/><w:keepNext/>'
            '<w:framePr w:dropCap="drop" w:lines="2" w:wrap="around" w:vAnchor="text" '
            'w:hAnchor="text"/>'
            '<w:spacing w:before="0" w:after="0" w:line="455" w:lineRule="exact"/>'
            '<w:ind w:firstLine="0"/><w:textAlignment w:val="baseline"/></w:pPr>'
            '<w:r><w:rPr><w:position w:val="-2"/><w:sz w:val="58"/></w:rPr>'
            f"<w:t>{letter}</w:t></w:r></w:p>"
        )
        p.addprevious(cap)
        set_ind(p, firstLine=0)

    def tall_math(self) -> None:
        """Exact line spacing, as LaTeX sets it, except where an equation needs the room."""
        tall = {
            f"{{{M_NS}}}{t}"
            for t in ("oMathPara", "f", "nary", "rad", "m", "eqArr", "limLow", "limUpp")
        }
        for p in self.body.iter(w("p")):
            if any(e.tag in tall for e in p.iter()):
                pp = ppr(p)
                sp = pp.find(w("spacing"))
                if sp is None:
                    sp = el("spacing")
                    pp.append(sp)
                sp.set(qn("w:line"), "230")
                sp.set(qn("w:lineRule"), "atLeast")

    def figures_paragraphs(self) -> None:
        for p in self.body.iter(w("p")):
            st = p.find(f"{w('pPr')}/{w('pStyle')}")
            if st is not None and st.get(qn("w:val")) == "CaptionedFigure":
                set_ind(p, firstLine=0)

    def tables(self) -> None:
        def rule(tag: str, sz: int) -> etree._Element:
            return el(tag, val="single", sz=sz, space=0, color="000000")

        for tbl in self.body.iter(w("tbl")):
            tp = tbl.find(w("tblPr"))
            for tag in ("tblBorders", "tblW", "jc", "tblCellMar", "tblLayout"):
                old = tp.find(w(tag))
                if old is not None:
                    tp.remove(old)
            tp.append(el("tblW", w=0, type="auto"))
            tp.append(el("jc", val="center"))
            borders = el("tblBorders")
            borders.append(rule("top", 8))
            borders.append(rule("bottom", 8))
            tp.append(borders)
            tp.append(el("tblLayout", type="autofit"))
            mar = el("tblCellMar")
            for side, v in (("top", 0), ("left", 45), ("bottom", 0), ("right", 45)):
                mar.append(el(side, w=v, type="dxa"))
            tp.append(mar)
            ordered_tblpr(tp)
            grid = tbl.find(w("tblGrid"))
            if grid is not None:
                # pandoc's widths, kept for tables whose p{} columns fix their proportions
                self.grids[tbl] = [int(gc.get(qn("w:w"), "0")) for gc in grid]
                for gc in grid:
                    gc.set(qn("w:w"), "1")
            rows = tbl.findall(w("tr"))
            for i, tr in enumerate(rows):
                for tc in tr.findall(w("tc")):
                    tcp = tc.find(w("tcPr"))
                    if tcp is None:
                        tcp = el("tcPr")
                        tc.insert(0, tcp)
                    old = tcp.find(w("tcW"))
                    if old is not None:
                        tcp.remove(old)
                    tcp.insert(0, el("tcW", w=0, type="auto"))
                    if i == 0:
                        old = tcp.find(w("tcBorders"))
                        if old is not None:
                            tcp.remove(old)
                        cb = el("tcBorders")
                        cb.append(rule("bottom", 4))
                        anchor = next(
                            x
                            for x in (tcp.find(w(t)) for t in ("vMerge", "gridSpan", "tcW"))
                            if x is not None
                        )
                        anchor.addnext(cb)
                    for p in tc.iter(w("p")):
                        set_style(p, "TableText")
                        for sz in list(p.iter(w("sz"), w("szCs"))):
                            sz.getparent().remove(sz)

    def lists(self) -> None:
        num = self.d.part.numbering_part.element
        for an in num.findall(w("abstractNum")):
            if int(an.get(qn("w:abstractNumId"))) in (900, 901):
                continue
            for lvl in an.findall(w("lvl")):
                ilvl = int(lvl.get(qn("w:ilvl")))
                fmt = lvl.find(w("numFmt"))
                text = lvl.find(w("lvlText"))
                if fmt is not None and fmt.get(qn("w:val")) == "decimal" and text is not None:
                    text.set(qn("w:val"), f"%{ilvl + 1})")
                lp = lvl.find(w("pPr"))
                if lp is None:
                    lp = el("pPr")
                    lvl.append(lp)
                for old in lp.findall(w("ind")):
                    lp.remove(old)
                lp.append(el("ind", left=360 + 360 * ilvl, hanging=260))

    def floats(self) -> None:
        # Take every float out of the text first, so that no float is anchored to a paragraph that
        # belongs to another (a text box inside a text box is something Word cannot open).
        taken: list[tuple[int, Float, list, etree._Element | None]] = []
        moved: set[int] = set()
        for fid, fl in enumerate(self.src.floats):
            b, e = self.markers.get(("FB", fid)), self.markers.get(("FE", fid))
            if b is None or e is None:
                self.warnings.append(f"float {fid} ({fl.label}) markers lost")
                continue
            content = []
            node = b.getnext()
            while node is not None and node is not e:
                content.append(node)
                node = node.getnext()
            moved.update(id(x) for x in (b, e, *content))
            taken.append((fid, fl, content, e))
        for fid, fl, content, e in taken:
            after = e.getnext()  # the fallback anchor: the next paragraph that stays in the text
            while after is not None and (after.tag != w("p") or id(after) in moved):
                after = after.getnext()
            taken[taken.index((fid, fl, content, e))] = (fid, fl, content, after)
        for fid, _fl, content, _a in taken:
            for x in (self.markers[("FB", fid)], *content, self.markers[("FE", fid)]):
                x.getparent().remove(x)
        boxes = []
        for _fid, fl, _content, _after in taken:
            box = self.pdf.float_box(fl) if self.pdf else None
            if box is None:
                self.warnings.append(f"{fl.label}: not found in the PDF; placed at the top")
                width = TEXT_W if fl.wide else COL_W
                box = Box(0, MARGIN_X, MARGIN_TOP, width, 150)
            boxes.append(box)
        # Anchors advance through the text as LaTeX's pages do: floats are placed in page order,
        # and each anchor is searched for only after the previous one.
        pages = sorted({b.page for b in boxes if b.float_page})
        runs: list[list[int]] = []
        for p in pages:
            if runs and p == runs[-1][-1] + 1:
                runs[-1].append(p)
            else:
                runs.append([p])
        run_of = {p: run for run in runs for p in run}
        holders: dict[int, etree._Element] = {}
        self.cursor = 0
        self.start_pages()
        order = sorted(range(len(taken)), key=lambda i: (boxes[i].page or 10**6, i))
        for i in order:
            _fid, fl, content, after = taken[i]
            box = boxes[i]
            if box.float_page:
                if box.page not in holders:
                    holders.update(self.float_page_holders(run_of[box.page], after))
                anchor = holders[box.page]
            else:
                found = self.anchor_paragraph(box.page) if box.page else None
                if found is None and box.page:
                    # nothing of the text is left to anchor to: a page of floats at the end
                    box.float_page = True
                    if box.page not in holders:
                        holders.update(self.float_page_holders([box.page], after))
                    found = holders[box.page]
                anchor = found if found is not None else after
            if "--debug-floats" in sys.argv:
                kids = list(self.body)
                print(
                    f"  {fl.label} p{box.page} fp={box.float_page} anchor#{kids.index(anchor)}"
                    f" {lead_text(anchor)[:50]!r}"
                )
            self.scale_images(content, box.w, fl.widths)
            for elem in content:
                for tbl in elem.iter(w("tbl")):
                    for side in ("left", "right"):
                        m = tbl.find(f"{w('tblPr')}/{w('tblCellMar')}/{w(side)}")
                        if m is not None:
                            m.set(qn("w:w"), str(int(fl.colsep * TW)))
                    for r in tbl.iter(w("r")):
                        rpr = r.find(w("rPr"))
                        if rpr is None:
                            rpr = el("rPr")
                            r.insert(0, rpr)
                        for old in rpr.findall(w("sz")) + rpr.findall(w("szCs")):
                            rpr.remove(old)
                        rpr.append(el("sz", val=int(fl.size * 2)))
                        rpr.append(el("szCs", val=int(fl.size * 2)))
                        ordered_rpr(rpr)
                    for p in tbl.iter(w("p")):
                        sp = ppr(p).find(w("spacing"))
                        if sp is None:
                            sp = el("spacing")
                            ppr(p).append(sp)
                        sp.set(qn("w:line"), str(int(fl.size * 1.2 * TW)))
                        sp.set(qn("w:lineRule"), "atLeast")
            for elem in content:  # notes under a table are set small, as in the PDF
                if elem.tag == w("p") and fl.kind == "table":
                    st = elem.find(f"{w('pPr')}/{w('pStyle')}")
                    if st is None or st.get(qn("w:val")) not in ("TableCaption", "TableLabel"):
                        set_style(elem, "TableNote")
            if not fl.nowrap:
                for elem in content:
                    for tbl in elem.iter(w("tbl")):
                        self.fixed_widths(tbl, fl.columns)
            if fl.nowrap:
                for elem in content:
                    for tcp in elem.iter(w("tcPr")):
                        before = [
                            tcp.find(w(t))
                            for t in ("shd", "tcBorders", "vMerge", "gridSpan", "tcW")
                        ]
                        next(x for x in before if x is not None).addnext(el("noWrap"))
            if content and content[-1].tag != w("p"):
                pad = el("p")
                pad.append(el("pPr"))
                set_style(pad, "FloatEnd")
                content.append(pad)
            run = self.textbox(content, box, top=box.y < PAGE_H / 2, wrap=not box.float_page)
            pp = anchor.find(w("pPr"))
            # after any float already anchored here, so floats keep their source (and SEQ) order
            pos = anchor.index(pp) + 1 if pp is not None else 0
            while (
                pos < len(anchor) and anchor[pos].find(f"{{{MC_NS}}}AlternateContent") is not None
            ):
                pos += 1
            anchor.insert(pos, run)

    def fixed_widths(self, tbl: etree._Element, columns: list[float | None]) -> None:
        """A table with p{} columns: those columns take LaTeX's width and wrap; the others keep
        their natural width and, as in LaTeX, never wrap."""
        if not columns:
            return
        for tr in tbl.findall(w("tr")):
            i = 0
            for tc in tr.findall(w("tc")):
                tcp = tc.find(w("tcPr"))
                span = tcp.find(w("gridSpan")) if tcp is not None else None
                n = int(span.get(qn("w:val"))) if span is not None else 1
                tcw = tcp.find(w("tcW")) if tcp is not None else None
                widths = columns[i : i + n]
                if tcw is not None and n == 1 and widths and widths[0] is not None:
                    tcw.set(qn("w:w"), str(int(widths[0] * TW)))
                    tcw.set(qn("w:type"), "dxa")
                    # LaTeX breaks a long path or URL in a p{} column (\nolinkurl); Word breaks
                    # only at spaces, so give code there invisible break points after / _ .
                    for r in tc.iter(w("r")):
                        st = r.find(f"{w('rPr')}/{w('rStyle')}")
                        t = r.find(w("t"))
                        if (
                            st is not None
                            and st.get(qn("w:val")) == "VerbatimChar"
                            and t is not None
                        ):
                            t.text = re.sub(r"([/_.])", "\\1\u200b", t.text or "")
                elif tcp is not None and widths and all(x is None for x in widths):
                    before = [
                        tcp.find(w(t)) for t in ("shd", "tcBorders", "vMerge", "gridSpan", "tcW")
                    ]
                    next(x for x in before if x is not None).addnext(el("noWrap"))
                i += n

    def float_page_holders(
        self, run: list[int], fallback: etree._Element
    ) -> dict[int, etree._Element]:
        """Pages LaTeX filled with floats become pages of their own in Word: for each, an empty
        paragraph that starts the page and holds the floats' anchors, then a page break, placed
        before the first paragraph LaTeX set after those pages. Their floats do not wrap text, so
        Word never pushes an anchor off its page (which would loop the layout forever)."""
        target = self.anchor_paragraph(run[-1] + 1)
        at_end = target is None  # LaTeX deferred these floats past the last paragraph
        if at_end:
            target = self.body.find(w("sectPr"))
            if target is None:
                target = fallback
        holders = {}
        for p in run if at_end else [*run, None]:
            para = el("p")
            para.append(el("pPr"))
            set_style(para, "FloatEnd")
            ppr(para).append(el("pageBreakBefore"))
            ordered_ppr(ppr(para))
            target.addprevious(para)
            if p is not None:
                holders[p] = para
        kids = list(self.body)
        self.cursor = len(kids) if at_end else kids.index(target)
        return holders

    def start_pages(self) -> None:
        """The PDF page on which each top-level paragraph begins, found by walking the document and
        the PDF's running text forward together."""
        self.starts: dict[etree._Element, int] = {}  # keyed by element: lxml reuses ids
        front = {"PaperTitle", "PaperAuthor", "AbstractText", "IndexTerms", "Separator", "Subtitle"}
        pos = 0
        for p in self.body.iterchildren(w("p")):
            st = p.find(f"{w('pPr')}/{w('pStyle')}")
            if st is not None and st.get(qn("w:val")) in front:
                continue  # the title block is not running text
            key = norm(lead_text(p))[:32]
            if len(key) < 16:
                continue
            # a paragraph starts within a page or so of the previous one; a match much further on
            # is the same words recurring later, not this paragraph
            hit = self.pdf.stream.find(key, pos, pos + 9000)
            if hit >= 0:
                self.starts[p] = self.pdf.page_of(hit)
                pos = hit + len(key)

    def anchor_paragraph(self, page: int) -> etree._Element | None:
        """The first paragraph, after the previous anchor, that begins on this PDF page; failing
        that (the page continues a paragraph begun earlier), the first that begins after it."""
        kids = list(self.body)
        later = None
        for i in range(self.cursor, len(kids)):
            start = self.starts.get(kids[i])
            if start is None:
                continue
            if start == page:
                self.cursor = i
                return kids[i]
            if start > page:
                later = i
                break
        if later is not None:
            self.cursor = later
            return kids[later]
        # No paragraph begins on or after this page (it only continues one): take the next paragraph
        # after the previous anchor, so that anchors, and the SEQ numbers that count them, stay
        # in order.
        for i in range(self.cursor + 1, len(kids)):
            if kids[i].tag == w("p"):
                self.cursor = i
                return kids[i]
        return None

    def scale_images(self, content: list, width_pt: float, fracs: list[float]) -> None:
        wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        i = 0
        for elem in content:
            for inline in elem.iter(f"{{{wp}}}inline"):
                ext = inline.find(f"{{{wp}}}extent")
                cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                frac = fracs[i] if i < len(fracs) else 1.0
                i += 1
                ncx = int(width_pt * frac * EMU)
                ncy = int(cy * ncx / cx)
                ext.set("cx", str(ncx))
                ext.set("cy", str(ncy))
                for x in inline.iter(f"{{{a}}}ext"):
                    x.set("cx", str(ncx))
                    x.set("cy", str(ncy))

    def textbox(self, content: list, box: Box, top: bool, wrap: bool = True) -> etree._Element:
        self.docpr += 1
        cx, cy = int(box.w * EMU), int(box.h * EMU)
        gap = int(11 * EMU)
        dist = f'distT="{0 if top else gap}" distB="{gap if top else 0}"'
        xml = f"""
<w:r xmlns:w="{W_NS}"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
 <mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>
  <wp:anchor {dist} distL="0" distR="0" simplePos="0" relativeHeight="{251660000 + self.docpr}"
     behindDoc="0" locked="0" layoutInCell="1" allowOverlap="0">
   <wp:simplePos x="0" y="0"/>
   <wp:positionH relativeFrom="page"><wp:posOffset>{int(box.x * EMU)}</wp:posOffset></wp:positionH>
   <wp:positionV relativeFrom="page"><wp:posOffset>{int(box.y * EMU)}</wp:posOffset></wp:positionV>
   <wp:extent cx="{cx}" cy="{cy}"/>
   <wp:effectExtent l="0" t="0" r="0" b="0"/>
   {"<wp:wrapTopAndBottom/>" if wrap else "<wp:wrapNone/>"}
   <wp:docPr id="{self.docpr}" name="Float {self.docpr}"/>
   <wp:cNvGraphicFramePr/>
   <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
    <wps:wsp>
     <wps:cNvSpPr txBox="1"/>
     <wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>
     <wps:txbx><w:txbxContent/></wps:txbx>
     <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"
        anchor="t" anchorCtr="0"><a:spAutoFit/></wps:bodyPr>
    </wps:wsp>
   </a:graphicData></a:graphic>
  </wp:anchor>
 </w:drawing></mc:Choice></mc:AlternateContent>
</w:r>"""
        run = parse_xml(re.sub(r">\s+<", "><", xml.strip()))
        holder = next(run.iter(w("txbxContent")))
        for elem in content:
            holder.append(elem)
        return run

    def sections(self) -> None:
        body = self.body
        final = body.find(w("sectPr"))
        if final is not None:
            body.remove(final)
        if self.layout.ieee:
            first = self.sect_pr(columns=1, continuous=False)
            last = self.sect_pr(columns=2, continuous=True)
            sep_p = next(
                (p for p in body.iter(w("p")) if p.find(f"{w('pPr')}/{w('pStyle')}") is not None
                 and p.find(f"{w('pPr')}/{w('pStyle')}").get(qn("w:val")) == "Separator"),
                None,
            )  # fmt: skip
            if sep_p is not None:
                ppr(sep_p).append(first)
            body.append(last)
        else:
            body.append(self.sect_pr(columns=1, continuous=False))
        self.headers()

    def sect_pr(self, columns: int, continuous: bool) -> etree._Element:
        s = el("sectPr")
        if continuous:
            s.append(el("type", val="continuous"))
        if self.layout.ieee:
            s.append(el("pgSz", w=PAGE_W * TW, h=PAGE_H * TW))
            s.append(
                el(
                    "pgMar",
                    top=MARGIN_TOP * TW,
                    right=MARGIN_X * TW,
                    bottom=MARGIN_BOTTOM * TW,
                    left=MARGIN_X * TW,
                    header=HEADER_DIST * TW,
                    footer=MARGIN_BOTTOM * TW // 2,
                    gutter=0,
                )
            )
            s.append(el("cols", num=columns, space=COL_GAP * TW))
        else:
            s.append(el("pgSz", w=PAGE_W * TW, h=PAGE_H * TW))
            s.append(
                el(
                    "pgMar",
                    top=1440,
                    right=1440,
                    bottom=1440,
                    left=1440,
                    header=720,
                    footer=720,
                    gutter=0,
                )
            )
            s.append(el("cols", space=720))
        return s

    def headers(self) -> None:
        sec = self.d.sections[0]
        if self.layout.ieee:
            hdr = sec.header
            hdr.is_linked_to_previous = False
            p = hdr.paragraphs[0]._p
            set_style(p, "RunningHead")
            p.append(text_run(self.layout.header.upper() + "\t"))
            for r in field_runs("PAGE", "1", None):
                p.append(r)
        else:
            ftr = sec.footer
            ftr.is_linked_to_previous = False
            p = ftr.paragraphs[0]._p
            set_style(p, "PageFoot")
            for r in field_runs("PAGE", "1", None):
                p.append(r)

    # -------------------------------------------------------------------------------- styles
    def styles(self) -> None:
        d = self.d
        ieee = self.layout.ieee
        st_el = d.styles.element
        for rf in st_el.iter(w("rFonts")):
            for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rf.attrib.pop(qn(f"w:{attr}"), None)
        for c in list(st_el.iter(w("color"))):
            c.getparent().remove(c)

        body_font = BODY_FONT if ieee else "Cambria"
        body_size = 9.5 if ieee else 10
        lead = 11.5 if ieee else 12.2

        def font(rpr_owner: object, name: str) -> None:
            rpr = rpr_owner.element.get_or_add_rPr()
            rf = rpr.find(w("rFonts"))
            if rf is None:
                rf = el("rFonts")
                rpr.insert(0, rf)
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn(f"w:{attr}"), name)

        def para(name: str, base: str | None = None, *, size: float | None = None,
                 face: str | None = None,
                 bold: bool | None = None, italic: bool | None = None, caps: bool = False,
                 align: WD_ALIGN_PARAGRAPH | None = None, before: float = 0, after: float = 0,
                 first: float | None = None, left: float | None = None, right: float | None = None,
                 line: float | None = None, exact: bool = True, keep: bool = False,
                 hyphen: bool = True) -> object:  # fmt: skip
            try:
                s = d.styles[name]
            except KeyError:
                s = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if base:
                s.base_style = d.styles[base]
            if face:
                font(s, face)
            f = s.font
            if size:
                f.size = Pt(size)
            if bold is not None:
                f.bold = bold
            if italic is not None:
                f.italic = italic
            f.small_caps = caps or None
            f.color.rgb = RGBColor(0, 0, 0)
            pf = s.paragraph_format
            if align is not None:
                pf.alignment = align
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            if first is not None:
                pf.first_line_indent = Pt(first)
            if left is not None:
                pf.left_indent = Pt(left)
            if right is not None:
                pf.right_indent = Pt(right)
            if line == -1:  # single spacing: a picture's height, not a text line's
                pf.line_spacing = 1.0
            elif line:
                pf.line_spacing = Pt(line)
                pf.line_spacing_rule = (
                    WD_LINE_SPACING.EXACTLY if exact else WD_LINE_SPACING.AT_LEAST
                )
            pf.keep_with_next = keep or None
            pf.widow_control = True
            if not hyphen:
                s.element.get_or_add_pPr().append(el("suppressAutoHyphens"))
            return s

        J, CTR, LEFT = (
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        )
        indent = 9.5 if ieee else 15
        para("Normal", face=body_font, size=body_size, align=J, line=lead, first=0)
        para("Body Text", "Normal", first=indent)
        para("First Paragraph", "Body Text", first=0)
        para("Compact", "Normal", first=0)
        para("Block Text", "Normal", left=indent, right=indent, first=0, before=3, after=3)
        if ieee:
            para("Heading 1", "Normal", face=SANS_FONT, size=11, bold=True, caps=True, align=LEFT,
                 before=14, after=5, line=13, keep=True, hyphen=False)  # fmt: skip
            para("Heading 2", "Normal", face=SANS_FONT, size=9.5, bold=True, align=LEFT, before=10,
                 after=4, keep=True, hyphen=False)  # fmt: skip
            para("Heading 3", "Normal", face=SANS_FONT, size=9.5, bold=True, italic=True,
                 align=LEFT, before=8, after=3, keep=True, hyphen=False)  # fmt: skip
        else:
            para("Heading 1", "Normal", face=body_font, size=14.4, bold=True, align=LEFT, before=18,
                 after=8, keep=True, hyphen=False)  # fmt: skip
            para("Heading 2", "Normal", face=body_font, size=12, bold=True, align=LEFT, before=12,
                 after=6, keep=True, hyphen=False)  # fmt: skip
        para("PaperTitle", "Normal", face=SANS_FONT if ieee else body_font,
             size=24 if ieee else 17.3,
             align=CTR, before=6 if ieee else 24, after=12 if ieee else 6, line=27 if ieee else 21,
             exact=True, hyphen=False)  # fmt: skip
        para("Subtitle", "Normal", size=12, align=CTR, after=30, line=14.5, hyphen=False)
        para("PaperAuthor", "Normal", face=SANS_FONT, size=11, align=CTR, after=20, hyphen=False)
        para("AbstractText", "Normal", face=SANS_FONT, size=8, left=20, right=20, line=9.6)
        para("IndexTerms", "AbstractText", before=7)
        para("Separator", "Normal", align=LEFT, left=150, before=8, after=20)
        para("TableLabel", "Normal", face=SANS_FONT, size=8, align=CTR, keep=True, hyphen=False)
        para(
            "Table Caption",
            "Normal",
            face=SANS_FONT,
            size=8,
            align=CTR,
            after=5,
            line=9.6,
            keep=True,
        )
        para("Image Caption", "Normal", face=SANS_FONT, size=8, align=J, before=5, line=9.6)
        para("Captioned Figure", "Normal", align=CTR, first=0, line=-1, keep=True)
        para("TableText", "Normal", size=8, align=LEFT, line=9.6, first=0, hyphen=False)
        para("FloatEnd", "Normal", size=1, line=1, exact=True)
        para("TableNote", "Normal", size=8, align=J, line=9.6, first=0, before=3)
        para("Reference", "Normal", size=8, align=J, line=9.4)
        para("Biography", "Normal", face=SANS_FONT, size=8, before=24, line=9.6)
        para("RunningHead", "Normal", face=SANS_FONT, size=7, align=LEFT, hyphen=False)
        para("PageFoot", "Normal", align=CTR)
        para("PointHead", "Normal", bold=True, align=LEFT, before=7, after=1, keep=True)
        para(
            "Footnote Text", "Normal", size=8, italic=ieee, align=LEFT, line=9.5, left=10, first=-10
        )
        if ieee:
            rh = d.styles["RunningHead"].element.get_or_add_pPr()
            tabs = el("tabs")
            tabs.append(el("tab", val="right", pos=TEXT_W * TW))
            rh.append(tabs)
            sp = d.styles["Separator"].element.get_or_add_pPr()
            tabs = el("tabs")
            tabs.append(el("tab", val="left", leader="underscore", pos=(TEXT_W // 2 - 5) * TW))
            tabs.append(el("tab", val="right", leader="underscore", pos=(TEXT_W - 150) * TW))
            sp.append(tabs)
        for name, face, size in (("Verbatim Char", MONO_FONT, None), ("Hyperlink", None, None)):
            try:
                s = d.styles[name]
            except KeyError:
                continue
            if face:
                font(s, face)
            if size:
                s.font.size = Pt(size)
            else:  # take the size of the text it sits in (8 pt in a table, 9.5 pt in the body)
                for sz in list(s.element.iter(w("sz"), w("szCs"))):
                    sz.getparent().remove(sz)
                based = s.element.find(w("basedOn"))
                if based is not None:  # pandoc bases it on Body Text Char, which fixes 9.5 pt
                    s.element.remove(based)
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.font.underline = False
        # Word's built-in defaults: no theme fonts, and the document's language for hyphenation.
        defaults = st_el.find(w("docDefaults"))
        if defaults is not None:
            for rf in defaults.iter(w("rFonts")):
                for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                    rf.set(qn(f"w:{attr}"), body_font)
            rpr = defaults.find(f"{w('rPrDefault')}/{w('rPr')}")
            if rpr is not None and rpr.find(w("lang")) is None:
                rpr.append(el("lang", val="en-US"))

    def numbering(self) -> None:
        num = self.d.part.numbering_part.element
        first_num = num.find(w("num"))

        def lvl(i: int, text: str, style: str | None, tab: int, fmt: str = "decimal") -> str:
            ps = f'<w:pStyle w:val="{style}"/>' if style else ""
            indent = tab if style is None else 0
            return (
                f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>{ps}'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/><w:suff w:val="tab"/>'
                f'<w:pPr><w:tabs><w:tab w:val="num" w:pos="{tab}"/></w:tabs>'
                f'<w:ind w:left="{indent}" w:hanging="{indent}"/>'
                f"</w:pPr></w:lvl>"
            )

        headings = (
            f'<w:abstractNum xmlns:w="{W_NS}" w:abstractNumId="900">'
            '<w:multiLevelType w:val="multilevel"/>'
            + lvl(0, "%1", "Heading1", 300)
            + lvl(1, "%1.%2", "Heading2", 470)
            + lvl(2, "%1.%2.%3", "Heading3", 600)
            + "</w:abstractNum>"
        )
        refs = (
            f'<w:abstractNum xmlns:w="{W_NS}" w:abstractNumId="901">'
            '<w:multiLevelType w:val="singleLevel"/>'
            + lvl(0, "[%1]", None, 380)
            + "</w:abstractNum>"
        )
        for xml in (headings, refs):
            an = parse_xml(xml)
            if first_num is not None:
                first_num.addprevious(an)
            else:
                num.append(an)
        for n in (900, 901):
            num.append(
                parse_xml(
                    f'<w:num xmlns:w="{W_NS}" w:numId="{n}"><w:abstractNumId w:val="{n}"/></w:num>'
                )
            )
        if self.layout.ieee:
            for i, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
                pp = self.d.styles[name].element.get_or_add_pPr()
                np_ = el("numPr")
                np_.append(el("ilvl", val=i))
                np_.append(el("numId", val=900))
                pp.insert(0, np_)
                ordered_ppr(pp)

    def settings(self) -> None:
        s = self.d.settings.element
        for tag in (
            "autoHyphenation",
            "hyphenationZone",
            "doNotHyphenateCaps",
            "consecutiveHyphenLimit",
        ):
            for old in s.findall(w(tag)):
                s.remove(old)
        extra = [el("autoHyphenation"), el("consecutiveHyphenLimit", val=3),
                 el("hyphenationZone", val=180),
                 el("doNotHyphenateCaps")]  # fmt: skip
        anchor = s.find(w("characterSpacingControl"))
        for e in reversed(extra):
            if anchor is not None:
                anchor.addprevious(e)
            else:
                s.insert(0, e)
        compat = s.find(w("compat"))
        if compat is None:
            compat = el("compat")
            s.append(compat)
        for cs in compat.findall(w("compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                compat.remove(cs)
        compat.append(
            el(
                "compatSetting",
                name="compatibilityMode",
                uri="http://schemas.microsoft.com/office/word",
                val=15,
            )
        )

    def footnote_mark(self) -> None:
        """The first-page author note takes a bullet, as IEEEtran's compsoc thanks do."""
        with zipfile.ZipFile(self.path) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        doc = etree.fromstring(parts["word/document.xml"])
        ids = []
        for fr in doc.iter(w("footnoteReference")):
            p = fr.getparent().getparent()
            st = p.find(f"{w('pPr')}/{w('pStyle')}")
            if st is not None and st.get(qn("w:val")) == "PaperAuthor":
                fr.set(qn("w:customMarkFollows"), "1")
                mark = text_run("", fr.getparent().find(w("rPr")))
                fr.getparent().addnext(mark)
                ids.append(fr.get(qn("w:id")))
        if not ids:
            return
        notes = etree.fromstring(parts["word/footnotes.xml"])
        for fn in notes.findall(w("footnote")):
            if fn.get(qn("w:id")) in ids:
                for ref in list(fn.iter(w("footnoteRef"))):
                    r = ref.getparent()
                    r.remove(ref)
                    t = el("t")
                    t.text = "\u2022"
                    r.append(t)
                    r.append(el("tab"))
        parts["word/document.xml"] = etree.tostring(
            doc, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        parts["word/footnotes.xml"] = etree.tostring(
            notes, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        tmp = self.path.with_suffix(".tmp")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            for n, data in parts.items():
                z.writestr(n, data)
        tmp.replace(self.path)


def ordered_rpr(rpr: etree._Element) -> None:
    order = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
             "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish",
             "webHidden",
             "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
             "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
             "specVanish", "oMath"]  # fmt: skip
    rank = {w(t): i for i, t in enumerate(order)}
    kids = sorted(rpr, key=lambda e: rank.get(e.tag, 99))
    for k in kids:
        rpr.remove(k)
        rpr.append(k)


def ordered_tblpr(tp: etree._Element) -> None:
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
             "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
             "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]  # fmt: skip
    rank = {w(t): i for i, t in enumerate(order)}
    kids = sorted(tp, key=lambda e: rank.get(e.tag, 99))
    for k in kids:
        tp.remove(k)
        tp.append(k)


def para_text_of(r: etree._Element) -> str:
    return "".join(t.text or "" for t in r.iter(w("t")))


# -------------------------------------------------------------------------------- Word itself

WORD_PASS = r"""
param([string]$Paths)
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  foreach ($p in $Paths.Split('|')) {
    $d = $word.Documents.Open($p, $false, $false, $false)
    foreach ($s in $d.Shapes) { if ($s.TextFrame.HasText) { $s.TextFrame.AutoSize = -1 } }
    $d.Fields.Update() | Out-Null
    foreach ($sec in $d.Sections) {
      foreach ($h in $sec.Headers) { $h.Range.Fields.Update() | Out-Null }
    }
    $d.Repaginate()
    "{0}: {1} pages" -f (Split-Path $p -Leaf), $d.ComputeStatistics(2)
    $d.Save()
    $d.Close()
  }
} finally { $word.Quit() }
"""


def word_pass(paths: list[Path]) -> None:
    """Let Word itself update every field and fit every text box, and save a Word-native file."""
    if sys.platform != "win32" or "--no-word" in sys.argv:
        return
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "word_pass.ps1"
        script.write_text(WORD_PASS, encoding="utf-8")
        args = [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-Paths",
        ]
        args.append("|".join(str(p) for p in paths))
        try:
            out = subprocess.run(args, capture_output=True, text=True, timeout=900)
        except (OSError, subprocess.TimeoutExpired) as exc:
            print(f"Word pass skipped: {exc}")
            return
        if out.returncode != 0:
            print("Word pass skipped (Word not available?):", out.stderr.strip()[-300:])
            return
        print(out.stdout.strip())


# --------------------------------------------------------------------------------------- main


def build(root: str, layout: Layout) -> Path:
    tex, pdf = PAPER / f"{root}.tex", PAPER / f"{root}.pdf"
    newest = max(
        p.stat().st_mtime for p in [tex, *PAPER.glob("sections/*.tex"), *PAPER.glob("tables/*.tex")]
    )
    if not pdf.exists() or pdf.stat().st_mtime < newest:
        sys.exit(f"{pdf.name} is older than its sources; build the paper first")
    doc, src = flatten(root, fields=root != "response")
    DIST.mkdir(parents=True, exist_ok=True)
    out = DIST / f"{root}.docx"
    if "--keep-flat" in sys.argv:
        (DIST / f"{root}.flat.tex").write_text(doc, encoding="utf-8")
    pypandoc.convert_text(
        doc, "docx", format="latex", outputfile=str(out),
        extra_args=[f"--resource-path={PAPER}", "--wrap=none"],
    )  # fmt: skip
    pdf_layout = PdfLayout(pdf) if layout.ieee else None
    b = Builder(out, src, layout, pdf_layout)
    b.run()
    for msg in b.warnings:
        print(f"  {root}: {msg}")
    return out


def main() -> None:
    outs = [
        build(
            "main", Layout(ieee=True, header="IEEE Transactions on Dependable and Secure Computing")
        ),
        build("supplement", Layout(ieee=True, header="Supplementary Material")),
        build("response", Layout(ieee=False)),
    ]
    word_pass(outs)
    for o in outs:
        print("wrote", o.relative_to(C.repo_root()))
    for stale in DIST.glob("*.tmp"):
        stale.unlink()


if __name__ == "__main__":
    main()
